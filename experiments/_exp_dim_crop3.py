"""零 overfit:裁切細化 — 把「數值表」用 PyMuPDF find_tables 單獨精準裁出(排除機構圖),
配上含圖的 anchor 裁切一起給 VLM。對照 crop2 80%。控制變因:prompt 與 crop2 完全相同,只加「表格精準裁切」。
"""
import os, sys, json, base64, re
sys.stdout.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)
os.environ["BACKEND"] = "azure_gpt4o"
from dotenv import load_dotenv
load_dotenv()
import fitz
import config
from llm_backend import get_backend
from reevaluate import cmp_numeric
import openpyxl
import crop_extractor

ORIG_KW = list(crop_extractor.CROP_ANCHOR_KEYWORDS)
backend = get_backend("azure_gpt4o")
DIM_FIELDS = [("Maximum Length (mm)", "L"), ("Maximum Width (mm)", "W"), ("Maximum Height (mm)", "H")]


def load_spec():
    wb = openpyxl.load_workbook(config.SPECBOOK_PATH)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        if rec.get("Part Number"):
            out[rec["Part Number"]] = rec
    return out


spec = load_spec()


def raster_full(pdf_path, zoom=5.0, maxp=6):
    doc = fitz.open(str(pdf_path))
    out = []
    mat = fitz.Matrix(zoom, zoom)
    for i in range(min(doc.page_count, maxp)):
        out.append(base64.b64encode(doc[i].get_pixmap(matrix=mat).tobytes("png")).decode("ascii"))
    doc.close()
    return out


def anchor_crops_and_pages(pdf):
    """回 (crops[(page1,b64)], anchor_page_idx_set)。原 anchor 失敗補 DIMENSIONS。"""
    crop_extractor.CROP_ANCHOR_KEYWORDS = ORIG_KW
    bb = crop_extractor.find_anchor_bboxes(pdf)
    if not bb:
        crop_extractor.CROP_ANCHOR_KEYWORDS = ORIG_KW + ["DIMENSIONS", "Dimensions"]
        bb = crop_extractor.find_anchor_bboxes(pdf)
    crops = crop_extractor.get_dimension_crops(pdf, zoom=6.0, max_crops=2)
    crop_extractor.CROP_ANCHOR_KEYWORDS = ORIG_KW
    pages = sorted({pi for pi, _ in bb})
    return crops, pages


def table_crops(pdf_path, page_idx, zoom=6.0, pad=6):
    """在指定頁 find_tables，裁切每個表格 bbox 高清。"""
    out = []
    doc = fitz.open(str(pdf_path))
    try:
        page = doc[page_idx]
        try:
            tabs = page.find_tables()
        except Exception:
            return out
        mat = fitz.Matrix(zoom, zoom)
        pr = page.rect
        for t in tabs.tables:
            x0, y0, x1, y1 = t.bbox
            clip = fitz.Rect(max(0, x0 - pad), max(0, y0 - pad),
                             min(pr.width, x1 + pad), min(pr.height, y1 + pad))
            try:
                pix = page.get_pixmap(matrix=mat, clip=clip)
                out.append(base64.b64encode(pix.tobytes("png")).decode("ascii"))
            except Exception:
                continue
    finally:
        doc.close()
    return out


PROMPT = """看 datasheet 截圖的 Package Outline / Mechanical Dimensions / DIMENSIONS 表
(不是 Land Pattern / Footprint / Suggested Solder Pad / Tape),抽這顆零件的「最大外形尺寸」。
給你的圖含:單獨裁切的數值表(讀數用)+ 含機構圖的區域(判斷長/寬/高軸向用)。

# 關鍵定義(通用,適用所有封裝)
- **Maximum Length / Maximum Width = 含引腳的最大外形尺寸(overall, lead-to-lead span)**,不是 body 本體尺寸。
  - 例:SOT-23 的 E(lead span ~2.5)是含引腳外形,E1(body ~1.3)是本體 → Width 取 E 那個含引腳的較大值。
  - 表中若同時有 body 符號與含引腳符號(如 E vs E1、或 D vs H_E / lead span),取「含引腳/外形較大」那個。
- **Maximum Height = 封裝總高**(常見符號 A,也可能標成 T 或 H)。
- 三個值都取該列的 **MAX 欄(最右/最大)**,不是 MIN / NOM / TYP。
- **若尺寸標成「中心值 ± 公差」格式(例如 1.55±0.1),Max = 中心值 + 公差(=1.65),不要只取中心值。**
- Length 取較長軸方向、Width 取較短軸方向的含引腳外形(依圖示方向)。
- 單位若為 inch,乘 25.4。

# 輸出 JSON
{{"Maximum Length (mm)": <數值或null>, "Maximum Width (mm)": <數值或null>, "Maximum Height (mm)": <數值或null>}}
只輸出 JSON。"""


def parse(t):
    t = t.strip()
    m = re.search(r"\{.*\}", t, re.DOTALL)
    return json.loads(m.group(0)) if m else {}


print("尺寸抽取 — 表格精準裁切 + 含圖裁切 + crop2 規則 — 全 10 顆\n", flush=True)
hit = 0
total = 0
for pn, pdf_fn in config.PART_TO_PDF.items():
    pdf = config.DATASHEETS_DIR / pdf_fn
    acrops, apages = anchor_crops_and_pages(pdf)
    tcrops = []
    for pi in apages[:2]:
        tcrops += table_crops(pdf, pi)
    imgs = tcrops[:2] + [b64 for _, b64 in acrops]
    src = f"tbl{len(tcrops[:2])}+anc{len(acrops)}"
    if not imgs:
        imgs = raster_full(pdf)
        src = "full(miss)"
    try:
        ans = parse(backend.call_multimodal(PROMPT, imgs, timeout=config.TIMEOUT))
    except Exception as ex:
        print(f"[{pn}] ({src}) FAIL: {ex}", flush=True)
        total += 3
        continue
    exp = spec[pn]
    parts = []
    for f, k in DIM_FIELDS:
        ok = cmp_numeric(ans.get(f), exp.get(f), config.NUMERIC_TOLERANCE)
        total += 1
        hit += int(ok)
        parts.append(f"{k}={'OK' if ok else 'X'}({ans.get(f)}/{exp.get(f)})")
    print(f"[{pn}] ({src}) {'  '.join(parts)}", flush=True)

print(f"\n尺寸格命中: {hit}/{total} = {hit/total*100:.0f}%  (V23 57% / overall 70% / 裁切 77% / crop2 80%)", flush=True)
