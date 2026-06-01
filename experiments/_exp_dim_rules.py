"""零 overfit:針對三類尺寸錯誤加通用規則(不綁 datasheet 名/數值,只給原則 + 叫模型重看)。
基於 crop2(裁切 crop+DIM,80%),只改 prompt 加強規則:
  - ±公差必須取 中心+公差
  - 同方向多符號取較大(含引腳 vs 本體)
  - 字母不代表軸 + 高度是扁平 SMD 三者中最小(合理性檢查)
  - 排除 Carrier Tape / Land Pattern 數字
測全 10 顆尺寸三欄。對照 crop2 80%(24/30)。
"""
import os, sys, json, base64, re
sys.stdout.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)
os.environ["BACKEND"] = "azure_gpt4o"
from dotenv import load_dotenv
load_dotenv()
import fitz
import openpyxl
import config
from llm_backend import get_backend
from reevaluate import cmp_numeric
import crop_extractor

ORIG_KW = list(crop_extractor.CROP_ANCHOR_KEYWORDS)
backend = get_backend("azure_gpt4o")
DIM_FIELDS = [("Maximum Length (mm)", "L"), ("Maximum Width (mm)", "W"), ("Maximum Height (mm)", "H")]


def load_spec():
    wb = openpyxl.load_workbook(config.SPECBOOK_PATH)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        if rec.get("Part Number"):
            out[rec["Part Number"]] = rec
    return out


spec = load_spec()


def raster_full(pdf_path, zoom=5.0, maxp=6):
    doc = fitz.open(str(pdf_path))
    out = []
    mat = fitz.Matrix(zoom, zoom)
    for i in range(min(doc.page_count, maxp)):
        out.append(base64.b64encode(doc[i].get_pixmap(matrix=mat).tobytes("png")).decode("ascii"))
    doc.close()
    return out


def crops_default(pdf):
    crop_extractor.CROP_ANCHOR_KEYWORDS = ORIG_KW
    return crop_extractor.get_dimension_crops(pdf, zoom=6.0, max_crops=2)


def crops_extended(pdf):
    crop_extractor.CROP_ANCHOR_KEYWORDS = ORIG_KW + ["DIMENSIONS", "Dimensions"]
    try:
        return crop_extractor.get_dimension_crops(pdf, zoom=6.0, max_crops=2)
    finally:
        crop_extractor.CROP_ANCHOR_KEYWORDS = ORIG_KW


PROMPT = """看 datasheet 截圖的 Package Outline / Mechanical Dimensions / DIMENSIONS 表,抽這顆零件的「最大外形尺寸」長/寬/高。

# 通用規則(適用所有封裝,每條都要檢查)
1. 三個值都取規格表的 Max 欄(最右/最大),不要取 Min / Nom / Typ。
2. 【±公差】尺寸若寫成「中心值 ± 公差」(例如 0.9±0.1、1.55±0.1),你回報的值必須是 中心值＋公差(0.9±0.1→1.0、1.55±0.1→1.65)。先算「中心＋公差」再報,絕不只報中心值。
3. 【含引腳取大】Length / Width 要的是「含引腳的最大外形」(lead-to-lead overall),不是本體。同一個方向若表上有多個候選符號(本體 vs 含引腳,例如 D 本體長 vs H 含引腳全長、E 本體 vs lead span),一律取「數值較大」的那個。
4. 【字母不代表軸】符號字母(A / D / E / H / L…)不代表它量哪個軸——不要看到「H」就當高度、看到「D」就當長度。要依圖上標註線的指向判斷它量的是哪個方向。
5. 【高度合理性檢查】這類扁平 SMD 封裝的高度(厚度)幾乎一定是長、寬、高三者中「最小」的。若你選出來的高度比長度或寬度還大,代表選錯了——回去重判哪個符號才是垂直厚度方向。
6. 【排除非本體表】只抽封裝本體尺寸表;不要抽 Carrier Tape(包裝載帶)、Land Pattern / Recommended Footprint(焊墊)的數字。
7. 單位若為 inch,乘 25.4。

# 輸出 JSON
{{"Maximum Length (mm)": <數值或null>, "Maximum Width (mm)": <數值或null>, "Maximum Height (mm)": <數值或null>}}
只輸出 JSON。"""


def parse(t):
    t = t.strip()
    m = re.search(r"\{.*\}", t, re.DOTALL)
    return json.loads(m.group(0)) if m else {}


print("尺寸抽取 — 裁切 + 四條通用規則強化(零 overfit)— 全 10 顆\n", flush=True)
hit = 0
total = 0
for pn, pdf_fn in config.PART_TO_PDF.items():
    pdf = config.DATASHEETS_DIR / pdf_fn
    crops = crops_default(pdf) or crops_extended(pdf)
    imgs = [b64 for _, b64 in crops] if crops else raster_full(pdf)
    src = f"crop{len(crops)}" if crops else "full"
    try:
        ans = parse(backend.call_multimodal(PROMPT, imgs, timeout=config.TIMEOUT))
    except Exception as ex:
        print(f"[{pn}] ({src}) FAIL: {ex}", flush=True)
        total += 3
        continue
    exp = spec[pn]
    parts = []
    for f, k in DIM_FIELDS:
        ok = cmp_numeric(ans.get(f), exp.get(f), config.NUMERIC_TOLERANCE)
        total += 1
        hit += int(ok)
        parts.append(f"{k}={'OK' if ok else 'X'}({ans.get(f)}/{exp.get(f)})")
    print(f"[{pn}] ({src}) {'  '.join(parts)}", flush=True)

print(f"\n尺寸格命中: {hit}/{total} = {hit/total*100:.0f}%  (crop2 規則前: 24/30)", flush=True)
