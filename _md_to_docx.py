"""把 REPORT.md 轉成 REPORT.docx（學術排版：封面 + 頁碼 + 章節分頁 + 字級層級 + 學術表格 + 行距 1.5）。
支援：# / ## / ### 標題、| 表格、**粗體**、`inline code`、> 引用、- 清單、--- 分隔、段落。
"""
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

SRC = Path(__file__).resolve().parent / "REPORT.md"
DST = Path(__file__).resolve().parent / "REPORT.docx"

CJK_FONT = "標楷體"
ASCII_FONT = "Times New Roman"
CODE_FONT = "Consolas"

# 封面內容
COVER_TITLE_LINE1 = "Datasheet 參數提取系統"
COVER_TITLE_LINE2 = "成果報告"
COVER_COURSE = "NTUST AI 課程　期末作業"
COVER_NAME = "黃姿晴"
COVER_DATE = "2026 / 05 / 30"


def set_fonts(run, ascii_font=ASCII_FONT, cjk_font=CJK_FONT):
    """中文用 cjk_font、英文/數字用 ascii_font（OOXML w:rFonts 控制）。"""
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cjk_font)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)


set_cjk = set_fonts


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
            set_fonts(r)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            set_fonts(r, ascii_font=CODE_FONT, cjk_font=CJK_FONT)
            r.font.color.rgb = RGBColor(0x8A, 0x2B, 0xE2)
        else:
            r = paragraph.add_run(part)
            set_fonts(r)


def split_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def is_table_sep(line):
    return bool(re.fullmatch(r"\s*\|?[\s:-]+\|[\s:|\-]*", line)) and "-" in line


def set_cell_shading(cell, color_hex):
    """設定 cell 背景色（hex 不含 #）。"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def set_cell_borders(cell, color="000000", sz="4"):
    """全 4 邊細黑線。"""
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:tcBorders"))
    if existing is not None:
        tcPr.remove(existing)
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_page_break(paragraph):
    """段落結尾插入 page break。"""
    run = paragraph.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._element.append(br)


def insert_page_field(paragraph):
    """段落內插入 PAGE 域（顯示當前頁碼）。"""
    run = paragraph.add_run()
    fldChar_begin = OxmlElement("w:fldChar")
    fldChar_begin.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.text = " PAGE "
    fldChar_end = OxmlElement("w:fldChar")
    fldChar_end.set(qn("w:fldCharType"), "end")
    run._element.append(fldChar_begin)
    run._element.append(instrText)
    run._element.append(fldChar_end)
    set_fonts(run)
    run.font.size = Pt(10)
    return run


def make_cover(doc):
    """封面頁：標題置中大字、課程、姓名、日期。"""
    # 上方留白約 1/4 頁
    for _ in range(4):
        doc.add_paragraph()

    # 主標題
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COVER_TITLE_LINE1)
    r.bold = True
    r.font.size = Pt(28)
    set_fonts(r)

    # 副標題
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COVER_TITLE_LINE2)
    r.bold = True
    r.font.size = Pt(24)
    set_fonts(r)

    # 中段留白
    for _ in range(8):
        doc.add_paragraph()

    # 課程
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COVER_COURSE)
    r.font.size = Pt(16)
    set_fonts(r)

    # 姓名
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COVER_NAME)
    r.font.size = Pt(16)
    set_fonts(r)

    # 日期
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COVER_DATE)
    r.font.size = Pt(14)
    set_fonts(r)


def setup_body_section(doc):
    """加 section break 進入本文：新頁開始、頁碼從 1、頁尾右下顯示「第 X 頁」。"""
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)

    # 頁碼從 1 起算
    sectPr = body_section._sectPr
    pgNumType = OxmlElement("w:pgNumType")
    pgNumType.set(qn("w:start"), "1")
    sectPr.append(pgNumType)

    # 把這個 section 的 footer unlink，避免跟封面共用
    body_section.footer.is_linked_to_previous = False

    # 頁尾：右下角「第 X 頁」
    footer_p = body_section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = footer_p.add_run("第 ")
    set_fonts(r1)
    r1.font.size = Pt(10)
    insert_page_field(footer_p)
    r2 = footer_p.add_run(" 頁")
    set_fonts(r2)
    r2.font.size = Pt(10)


def apply_heading_style(heading, level):
    """H1 18pt 粗置中、H2 14pt 粗、H3 12pt 粗、H4 11pt 粗。"""
    sizes = {1: 18, 2: 14, 3: 12, 4: 11}
    sz = sizes.get(level, 11)
    if level == 1:
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # heading 前後加點間距
    heading.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.line_spacing = 1.2
    for r in heading.runs:
        r.bold = True
        r.font.size = Pt(sz)
        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        set_fonts(r)


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()

    # Normal 樣式：標楷體 + Times New Roman 12pt + 行距 1.5
    normal = doc.styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = Pt(12)
    rfonts = normal.element.rPr.rFonts
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    rfonts.set(qn("w:ascii"), ASCII_FONT)
    rfonts.set(qn("w:hAnsi"), ASCII_FONT)
    normal.paragraph_format.line_spacing = 1.5

    # 封面頁（section 1）
    make_cover(doc)

    # section break 進入本文 + 頁碼
    setup_body_section(doc)

    # 跳過 markdown 第一個 H1（已放封面）+ 緊接的 metadata 行
    skipped_title = False

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 空行
        if not stripped:
            i += 1
            continue

        # 分隔線
        if re.fullmatch(r"-{3,}", stripped):
            i += 1
            continue

        # 標題
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)

            # 跳過 markdown 第一個 H1（標題）+ 下一個 metadata 行
            if level == 1 and not skipped_title:
                skipped_title = True
                i += 1
                while i < n and not lines[i].strip():
                    i += 1
                # 課程/姓名/日期那行也跳過
                if i < n and ("NTUST" in lines[i] or "黃姿晴" in lines[i]):
                    i += 1
                continue

            # 真正章節 H1 — 開新頁
            if level == 1:
                pb_p = doc.add_paragraph()
                pb_p.paragraph_format.line_spacing = 1.0
                add_page_break(pb_p)

            h = doc.add_heading(level=min(level, 4))
            add_inline(h, text)
            apply_heading_style(h, level)
            i += 1
            continue

        # 引用塊
        if stripped.startswith(">"):
            quote_lines = []
            while i < n and lines[i].strip().startswith(">"):
                q = lines[i].strip().lstrip(">").strip()
                if q:
                    quote_lines.append(q)
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.right_indent = Cm(0.4)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, "　".join(quote_lines))
            for r in p.runs:
                r.italic = True
            continue

        # 表格
        if stripped.startswith("|"):
            tbl_lines = []
            while i < n and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            header = split_row(tbl_lines[0])
            body = [split_row(r) for r in tbl_lines[1:] if not is_table_sep(r)]
            ncol = len(header)
            table = doc.add_table(rows=1, cols=ncol)
            table.autofit = True

            # 表頭：灰底、粗體、邊框
            for j, htxt in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.paragraphs[0].text = ""
                cell.paragraphs[0].paragraph_format.line_spacing = 1.0
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                add_inline(cell.paragraphs[0], htxt)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
                set_cell_shading(cell, "D9D9D9")
                set_cell_borders(cell)

            for row in body:
                cells = table.add_row().cells
                for j in range(ncol):
                    txt = row[j] if j < len(row) else ""
                    cells[j].paragraphs[0].text = ""
                    cells[j].paragraphs[0].paragraph_format.line_spacing = 1.0
                    cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
                    add_inline(cells[j].paragraphs[0], txt)
                    set_cell_borders(cells[j])

            # 表格後留小間距
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(0)
            continue

        # 清單
        if re.match(r"^[-*]\s+", stripped):
            while i < n and re.match(r"^[-*]\s+", lines[i].strip()):
                item = re.sub(r"^[-*]\s+", "", lines[i].strip())
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, item)
                i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            while i < n and re.match(r"^\d+\.\s+", lines[i].strip()):
                item = re.sub(r"^\d+\.\s+", "", lines[i].strip())
                p = doc.add_paragraph(style="List Number")
                add_inline(p, item)
                i += 1
            continue

        # 一般段落
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(DST)
    print(f"Wrote: {DST}")
    print(f"段落: {len(doc.paragraphs)} / 表格: {len(doc.tables)} / sections: {len(doc.sections)}")


if __name__ == "__main__":
    main()
