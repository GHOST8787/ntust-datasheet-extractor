"""把 REPORT_formal.md 轉成正式排版 docx，套用學術論文寫作格式：
- 標題標楷體 / 內文新細明體 12 / 英文 Times New Roman 12 / inline code Consolas
- 封面（國立臺灣科技大學 + 課名 + 報告題目 + 姓名 + 日期，置中）
- 三段式 section：封面（無頁碼）/ 目錄頁（羅馬數字置中）/ 本文（阿拉伯數字置中，從 1 起算）
- 目錄（章 + 節，Word TOC 域）+ 表目次（依表標題樣式，Word TOC 域）
- 標題層級 章 / 節 / 壹一（一）；表標題在表上方、新細明體粗體
- 行距 1.5、表格內單行距、學術表格（表頭灰底 + 細黑線）
- 邊界 上3 下2 左3 右2.5 cm
* 不含碩論專屬件（口試審定書 / 指導教授 / 中英文摘要），身分為台科大課程作業。
"""
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE

import argparse

_HERE = Path(__file__).resolve().parent
_ap = argparse.ArgumentParser(description="REPORT_formal.md → 學術排版 docx（可參數化封面，預設為 DATASHEET 報告）")
_ap.add_argument("--src", default=str(_HERE / "REPORT_formal.md"))
_ap.add_argument("--dst", default=str(_HERE / "20260602_黃姿晴_NTUST_DATASHEET_期末報告.docx"))
_ap.add_argument("--school", default="國立臺灣科技大學")
_ap.add_argument("--course", default="NTUST_DATASHEET　期末報告")
_ap.add_argument("--title1", default="Datasheet 參數提取系統")
_ap.add_argument("--title2", default="成果報告")
_ap.add_argument("--name", default="黃姿晴")
_ap.add_argument("--date", default="中華民國 115 年 6 月")
_args = _ap.parse_args()

SRC = Path(_args.src)
DST = Path(_args.dst)

CJK_BODY = "新細明體"       # 內文
CJK_HEADING = "標楷體"      # 標題
ASCII_FONT = "Times New Roman"
CODE_FONT = "Consolas"

# 封面內容（台科大身分，非碩論；可由 argv 覆蓋）
COVER_SCHOOL = _args.school
COVER_COURSE = _args.course
COVER_TITLE_LINE1 = _args.title1
COVER_TITLE_LINE2 = _args.title2
COVER_NAME = _args.name
COVER_DATE = _args.date

# 表標題 / 附表標題 偵測
CAP_RE = re.compile(r"^(表|附表)\s*\S")


def set_fonts(run, ascii_font=ASCII_FONT, cjk_font=CJK_BODY):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cjk_font)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline(paragraph, text, cjk_font=CJK_BODY):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
            set_fonts(r, cjk_font=cjk_font)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            set_fonts(r, ascii_font=CODE_FONT, cjk_font=cjk_font)
            r.font.color.rgb = RGBColor(0x8A, 0x2B, 0xE2)
        else:
            r = paragraph.add_run(part)
            set_fonts(r, cjk_font=cjk_font)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_table_sep(line):
    return bool(re.fullmatch(r"\s*\|?[\s:-]+\|[\s:|\-]*", line)) and "-" in line


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def set_cell_borders(cell, color="000000", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:tcBorders"))
    if existing is not None:
        tcPr.remove(existing)
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), sz); el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_page_break(paragraph):
    run = paragraph.add_run()
    br = OxmlElement("w:br"); br.set(qn("w:type"), "page")
    run._element.append(br)


def insert_page_field(paragraph, size=12, cjk=CJK_BODY):
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._element.append(b); run._element.append(it); run._element.append(e)
    set_fonts(run, cjk_font=cjk); run.font.size = Pt(size)


def add_toc_field(doc, switches, placeholder):
    """插入 Word TOC 域（開檔後按 F9 或 Ctrl+A→F9 更新頁碼）。"""
    p = doc.add_paragraph()
    r1 = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = switches
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    r1._element.append(b); r1._element.append(it); r1._element.append(sep)
    r2 = p.add_run(placeholder); set_fonts(r2)
    r3 = p.add_run()
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r3._element.append(e)
    return p


def set_section_pagenum(section, fmt=None, start=None):
    sectPr = section._sectPr
    pg = OxmlElement("w:pgNumType")
    if fmt:
        pg.set(qn("w:fmt"), fmt)        # upperRoman / decimal
    if start:
        pg.set(qn("w:start"), str(start))
    sectPr.append(pg)


def set_centered_footer(section, fmt_roman=False):
    section.footer.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    insert_page_field(fp)


def set_margins(section):
    section.top_margin = Cm(3); section.bottom_margin = Cm(2)
    section.left_margin = Cm(3); section.right_margin = Cm(2.5)


def make_cover(doc):
    """封面：單行距 + 精簡空白，確保整頁不溢出到第二頁。"""
    def blank(n=1):
        for _ in range(n):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
    def center(text, size, bold=True, cjk=CJK_HEADING, ascii_f=ASCII_FONT):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
        set_fonts(r, ascii_font=ascii_f, cjk_font=cjk)
        return p
    blank(2)
    center(COVER_SCHOOL, 22)
    center(COVER_COURSE, 16)
    blank(4)
    center(COVER_TITLE_LINE1, 28)
    center(COVER_TITLE_LINE2, 24)
    blank(5)
    center(COVER_NAME, 18)
    center(COVER_DATE, 16)


def add_section_title(doc, text):
    """目錄 / 表目次 之類的置中大標（標楷體 22 粗）。"""
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12); p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text); r.bold = True; r.font.size = Pt(22)
    set_fonts(r, cjk_font=CJK_HEADING)
    return p


def apply_heading_style(heading, level):
    """章=18 置中 / 節=15 / 壹=13 / 一=12，皆標楷體粗。"""
    sizes = {1: 18, 2: 15, 3: 13, 4: 12}
    sz = sizes.get(level, 12)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    heading.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    heading.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    heading.paragraph_format.line_spacing = 1.3
    for r in heading.runs:
        r.bold = True; r.font.size = Pt(sz); r.font.color.rgb = RGBColor(0, 0, 0)
        set_fonts(r, cjk_font=CJK_HEADING)


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()

    # Normal：新細明體 + Times New Roman 12 + 行距 1.5
    normal = doc.styles["Normal"]
    normal.font.name = ASCII_FONT; normal.font.size = Pt(12)
    rfonts = normal.element.rPr.rFonts
    rfonts.set(qn("w:eastAsia"), CJK_BODY)
    rfonts.set(qn("w:ascii"), ASCII_FONT); rfonts.set(qn("w:hAnsi"), ASCII_FONT)
    normal.paragraph_format.line_spacing = 1.5

    # 表標題樣式（給表目次 TOC 收集用）：新細明體 12 粗
    cap_style = doc.styles.add_style("TableCaption", WD_STYLE_TYPE.PARAGRAPH)
    cap_style.font.name = ASCII_FONT; cap_style.font.size = Pt(12); cap_style.font.bold = True
    cap_rfonts = cap_style.element.get_or_add_rPr().get_or_add_rFonts()
    cap_rfonts.set(qn("w:eastAsia"), CJK_BODY)
    cap_rfonts.set(qn("w:ascii"), ASCII_FONT); cap_rfonts.set(qn("w:hAnsi"), ASCII_FONT)
    cap_style.paragraph_format.line_spacing = 1.2
    cap_style.paragraph_format.space_before = Pt(6); cap_style.paragraph_format.space_after = Pt(2)

    # 目錄/表目次條目（Word 內建 TOC 1~4 樣式）明確指定新細明體，
    # 不再依賴佈景預設字型（解決「目錄條目字型空白/不確定」問題）
    for _tname in ("TOC 1", "TOC 2", "TOC 3", "TOC 4"):
        try:
            _tst = doc.styles[_tname]
        except KeyError:
            _tst = doc.styles.add_style(_tname, WD_STYLE_TYPE.PARAGRAPH)
        _trpr = _tst.element.get_or_add_rPr()
        _trf = _trpr.find(qn("w:rFonts"))
        if _trf is None:
            _trf = OxmlElement("w:rFonts"); _trpr.append(_trf)
        _trf.set(qn("w:eastAsia"), CJK_BODY)
        _trf.set(qn("w:ascii"), ASCII_FONT)
        _trf.set(qn("w:hAnsi"), ASCII_FONT)
        _tst.font.size = Pt(12)

    # ===== section 1：封面 =====
    make_cover(doc)

    # ===== section 2：目錄頁（羅馬數字置中）=====
    toc_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_pagenum(toc_section, fmt="upperRoman", start=1)
    set_centered_footer(toc_section)

    add_section_title(doc, "目　錄")
    add_toc_field(doc, 'TOC \\o "1-2" \\h \\z \\u', "（請在 Word 內按 Ctrl+A 後 F9 更新目錄頁碼）")

    # 表目次另起一頁
    pb = doc.add_paragraph(); pb.paragraph_format.line_spacing = 1.0; add_page_break(pb)
    add_section_title(doc, "表目次")
    add_toc_field(doc, 'TOC \\h \\z \\t "TableCaption,1"', "（請在 Word 內按 F9 更新表目次頁碼）")

    # ===== section 3：本文（阿拉伯數字置中，從 1 起算）=====
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_pagenum(body_section, fmt="decimal", start=1)
    set_centered_footer(body_section)

    # 跳過 markdown 開頭的標題/metadata 純文字行，從第一個 # 標題開始
    i = 0; n = len(lines)
    while i < n and not re.match(r"^#{1,6}\s", lines[i].strip()):
        i += 1

    first_chapter = True
    while i < n:
        stripped = lines[i].strip()
        if not stripped:
            i += 1; continue
        if re.fullmatch(r"-{3,}", stripped):
            i += 1; continue

        # 標題
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1)); text = m.group(2)
            if level == 1 and not first_chapter:
                pbp = doc.add_paragraph(); pbp.paragraph_format.line_spacing = 1.0
                add_page_break(pbp)
            if level == 1:
                first_chapter = False
            h = doc.add_heading(level=min(level, 4))
            add_inline(h, text, cjk_font=CJK_HEADING)
            apply_heading_style(h, level)
            i += 1; continue

        # 引用塊（>4 行內縮斜體 10）
        if stripped.startswith(">"):
            qlines = []
            while i < n and lines[i].strip().startswith(">"):
                q = lines[i].strip().lstrip(">").strip()
                if q:
                    qlines.append(q)
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0); p.paragraph_format.right_indent = Cm(0)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, "　".join(qlines))
            for r in p.runs:
                r.italic = True; r.font.size = Pt(10)
            continue

        # 表標題（表 X-Y… / 附表 N…）→ TableCaption 樣式，置於表上方
        if CAP_RE.match(stripped):
            p = doc.add_paragraph(style="TableCaption")
            add_inline(p, stripped, cjk_font=CJK_BODY)
            i += 1; continue

        # 表格
        if stripped.startswith("|"):
            tbl = []
            while i < n and lines[i].strip().startswith("|"):
                tbl.append(lines[i]); i += 1
            header = split_row(tbl[0])
            body = [split_row(r) for r in tbl[1:] if not is_table_sep(r)]
            ncol = len(header)
            table = doc.add_table(rows=1, cols=ncol); table.autofit = True
            for j, htxt in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.paragraphs[0].text = ""
                cell.paragraphs[0].paragraph_format.line_spacing = 1.0
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                add_inline(cell.paragraphs[0], htxt, cjk_font=CJK_BODY)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
                set_cell_shading(cell, "D9D9D9"); set_cell_borders(cell)
            for row in body:
                cells = table.add_row().cells
                for j in range(ncol):
                    txt = row[j] if j < len(row) else ""
                    cells[j].paragraphs[0].text = ""
                    cells[j].paragraphs[0].paragraph_format.line_spacing = 1.0
                    cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
                    add_inline(cells[j].paragraphs[0], txt, cjk_font=CJK_BODY)
                    set_cell_borders(cells[j])
            spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(0)
            continue

        # 清單
        if re.match(r"^[-*]\s+", stripped):
            while i < n and re.match(r"^[-*]\s+", lines[i].strip()):
                item = re.sub(r"^[-*]\s+", "", lines[i].strip())
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, item)
                i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            while i < n and re.match(r"^\d+\.\s+", lines[i].strip()):
                item = re.sub(r"^\d+\.\s+", "", lines[i].strip())
                p = doc.add_paragraph(style="List Number")
                add_inline(p, item)
                i += 1
            continue

        # 一般段落（首行縮排 2 字）
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(24)
        add_inline(p, stripped)
        i += 1

    # 全 section 邊界
    for s in doc.sections:
        set_margins(s)

    doc.save(DST)
    print(f"Wrote: {DST}")
    print(f"段落 {len(doc.paragraphs)} / 表格 {len(doc.tables)} / sections {len(doc.sections)}")


if __name__ == "__main__":
    main()
